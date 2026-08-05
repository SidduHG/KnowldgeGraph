from pathlib import Path
from PIL import Image, ImageDraw, ImageFont

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt


DESKTOP = Path(r"C:\Users\siddu\OneDrive\Desktop\Desktop")
SRC = DESKTOP / "final ppt.pptx"
OUT = DESKTOP / "final ppt_AURA_completed.pptx"

NAVY = RGBColor(24, 45, 74)
TEAL = RGBColor(15, 110, 86)
GOLD = RGBColor(186, 117, 23)
ROSE = RGBColor(153, 53, 86)
SLATE = RGBColor(92, 104, 120)
LIGHT = RGBColor(244, 248, 250)
INK = RGBColor(28, 35, 45)
WHITE = RGBColor(255, 255, 255)


def font(size=18, bold=False, color=INK):
    return {"size": Pt(size), "bold": bold, "color": color}


def clear_slide(slide):
    for shape in list(slide.shapes):
        shape.element.getparent().remove(shape.element)


def set_text(shape, text, size=18, bold=False, color=INK, align=PP_ALIGN.LEFT):
    shape.text = text
    tf = shape.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.08)
    tf.margin_right = Inches(0.08)
    tf.margin_top = Inches(0.04)
    tf.margin_bottom = Inches(0.04)
    for p in tf.paragraphs:
        p.alignment = align
        for r in p.runs:
            r.font.name = "Aptos"
            r.font.size = Pt(size)
            r.font.bold = bold
            r.font.color.rgb = color


def add_box(slide, x, y, w, h, text, fill=LIGHT, line=TEAL, size=15, bold=True, color=INK):
    shp = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, x, y, w, h)
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    shp.line.color.rgb = line
    shp.line.width = Pt(1.5)
    shp.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    set_text(shp, text, size=size, bold=bold, color=color, align=PP_ALIGN.CENTER)
    return shp


def add_textbox(slide, x, y, w, h, text, size=18, bold=False, color=INK, align=PP_ALIGN.LEFT):
    shp = slide.shapes.add_textbox(x, y, w, h)
    set_text(shp, text, size=size, bold=bold, color=color, align=align)
    return shp


def add_arrow(slide, x1, y1, x2, y2, color=SLATE, width=2.0):
    conn = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, x1, y1, x2, y2)
    conn.line.color.rgb = color
    conn.line.width = Pt(width)
    conn.line.end_arrowhead = True
    return conn


def add_header(slide, title, num):
    add_textbox(slide, Inches(0.45), Inches(0.22), Inches(11.6), Inches(0.45), title, 25, True, NAVY)
    line = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0.45), Inches(0.78), Inches(12.4), Inches(0.03))
    line.fill.solid()
    line.fill.fore_color.rgb = TEAL
    line.line.fill.background()
    add_textbox(slide, Inches(0.45), Inches(7.08), Inches(6.8), Inches(0.2), "Department of Information Science & Engineering, DSCE", 8.5, False, SLATE)
    add_textbox(slide, Inches(12.45), Inches(7.08), Inches(0.4), Inches(0.2), str(num), 8.5, False, SLATE, PP_ALIGN.RIGHT)


def add_bullets(slide, items, x, y, w, h, size=18, color=INK):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.1)
    tf.margin_right = Inches(0.1)
    tf.margin_top = Inches(0.04)
    tf.margin_bottom = Inches(0.04)
    tf.clear()
    for idx, item in enumerate(items):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.text = item
        p.level = 0
        p.font.name = "Aptos"
        p.font.size = Pt(size)
        p.font.color.rgb = color
        p.space_after = Pt(7)
    return box


def draw_png_diagram(path: Path, title: str, nodes, arrows):
    W, H = 1600, 900
    img = Image.new("RGB", (W, H), (248, 251, 252))
    d = ImageDraw.Draw(img)
    try:
        title_font = ImageFont.truetype("arialbd.ttf", 42)
        node_font = ImageFont.truetype("arialbd.ttf", 24)
        small_font = ImageFont.truetype("arial.ttf", 20)
    except Exception:
        title_font = node_font = small_font = ImageFont.load_default()
    d.text((60, 40), title, fill=(24, 45, 74), font=title_font)
    d.line((60, 105, 1540, 105), fill=(15, 110, 86), width=5)
    def edge_point(src, dst):
        sx, sy, sw, sh, *_ = src
        dx, dy, dw, dh, *_ = dst
        cx, cy = sx + sw / 2, sy + sh / 2
        tx, ty = dx + dw / 2, dy + dh / 2
        vx, vy = tx - cx, ty - cy
        if abs(vx) * sh > abs(vy) * sw:
            ex = cx + (sw / 2) * (1 if vx > 0 else -1)
            ey = cy + vy * ((sw / 2) / abs(vx))
        else:
            ey = cy + (sh / 2) * (1 if vy > 0 else -1)
            ex = cx + vx * ((sh / 2) / abs(vy)) if vy else cx
        return ex, ey

    for a, b in arrows:
        ax, ay, aw, ah, *_ = nodes[a]
        bx, by, bw, bh, *_ = nodes[b]
        x1, y1 = edge_point(nodes[a], nodes[b])
        x2, y2 = edge_point(nodes[b], nodes[a])
        d.line((x1, y1, x2, y2), fill=(92, 104, 120), width=4)
        # arrow head
        import math
        ang = math.atan2(y2 - y1, x2 - x1)
        r = 18
        p1 = (x2 - r * math.cos(ang - 0.45), y2 - r * math.sin(ang - 0.45))
        p2 = (x2 - r * math.cos(ang + 0.45), y2 - r * math.sin(ang + 0.45))
        d.polygon([(x2, y2), p1, p2], fill=(92, 104, 120))
    for key, (x, y, w, h, text, fill, outline) in nodes.items():
        d.rounded_rectangle((x, y, x + w, y + h), radius=22, fill=fill, outline=outline, width=4)
        lines = text.split("\n")
        total_h = len(lines) * 28
        ty = y + (h - total_h) // 2
        for line in lines:
            bbox = d.textbbox((0, 0), line, font=node_font if len(line) < 18 else small_font)
            d.text((x + (w - (bbox[2] - bbox[0])) / 2, ty), line, fill=(28, 35, 45), font=node_font if len(line) < 18 else small_font)
            ty += 30
    img.save(path)


def draw_sequence_png(path: Path):
    W, H = 1600, 900
    img = Image.new("RGB", (W, H), (248, 251, 252))
    d = ImageDraw.Draw(img)
    try:
        title_font = ImageFont.truetype("arialbd.ttf", 42)
        head_font = ImageFont.truetype("arialbd.ttf", 23)
        msg_font = ImageFont.truetype("arial.ttf", 20)
    except Exception:
        title_font = head_font = msg_font = ImageFont.load_default()
    d.text((60, 40), "AURA Sequence Diagram", fill=(24, 45, 74), font=title_font)
    d.line((60, 105, 1540, 105), fill=(15, 110, 86), width=5)
    actors = [("User", 170), ("React UI", 430), ("FastAPI", 690), ("AI/LLM Layer", 960), ("DB + Reports", 1260)]
    for name, x in actors:
        d.rounded_rectangle((x - 105, 145, x + 105, 215), radius=18, fill=(231, 241, 252), outline=(24, 95, 165), width=4)
        bb = d.textbbox((0, 0), name, font=head_font)
        d.text((x - (bb[2] - bb[0]) / 2, 168), name, fill=(28, 35, 45), font=head_font)
        d.line((x, 215, x, 785), fill=(110, 120, 135), width=3)
    xmap = dict(actors)
    steps = [
        ("Submit query / upload data", "User", "React UI"),
        ("POST request", "React UI", "FastAPI"),
        ("Preprocess + schema prompt", "FastAPI", "AI/LLM Layer"),
        ("Return schema-aware SQL", "AI/LLM Layer", "FastAPI"),
        ("Execute validated SQL", "FastAPI", "DB + Reports"),
        ("Records + generated report", "DB + Reports", "FastAPI"),
        ("Response + download link", "FastAPI", "React UI"),
        ("Display report/dashboard", "React UI", "User"),
    ]
    y = 275
    for label, src, dst in steps:
        x1, x2 = xmap[src], xmap[dst]
        d.line((x1, y, x2, y), fill=(15, 110, 86), width=4)
        import math
        ang = 0 if x2 > x1 else math.pi
        p1 = (x2 - 16 * math.cos(ang - 0.45), y - 16 * math.sin(ang - 0.45))
        p2 = (x2 - 16 * math.cos(ang + 0.45), y - 16 * math.sin(ang + 0.45))
        d.polygon([(x2, y), p1, p2], fill=(15, 110, 86))
        bb = d.textbbox((0, 0), label, font=msg_font)
        d.rectangle((min(x1, x2) + 12, y - 31, min(x1, x2) + 24 + (bb[2] - bb[0]), y - 6), fill=(248, 251, 252))
        d.text((min(x1, x2) + 18, y - 31), label, fill=(28, 35, 45), font=msg_font)
        y += 62
    img.save(path)


def draw_architecture_png(path: Path):
    W, H = 2400, 1350
    img = Image.new("RGB", (W, H), (246, 250, 252))
    d = ImageDraw.Draw(img)
    try:
        title_font = ImageFont.truetype("arialbd.ttf", 62)
        h_font = ImageFont.truetype("arialbd.ttf", 28)
        node_font = ImageFont.truetype("arialbd.ttf", 17)
        small_font = ImageFont.truetype("arial.ttf", 15)
        tiny_font = ImageFont.truetype("arial.ttf", 13)
        icon_font = ImageFont.truetype("seguisym.ttf", 26)
    except Exception:
        title_font = h_font = node_font = small_font = tiny_font = icon_font = ImageFont.load_default()

    navy = (24, 45, 74)
    teal = (15, 110, 86)
    blue = (24, 95, 165)
    gold = (186, 117, 23)
    rose = (153, 53, 86)
    slate = (92, 104, 120)
    ink = (28, 35, 45)
    grid = (221, 229, 235)

    def text_center(box, text, font, fill=ink, line_gap=4):
        x, y, w, h = box
        lines = text.split("\n")
        heights = [d.textbbox((0, 0), line, font=font)[3] for line in lines]
        total = sum(heights) + line_gap * (len(lines) - 1)
        cy = y + (h - total) / 2
        for line, lh in zip(lines, heights):
            bb = d.textbbox((0, 0), line, font=font)
            d.text((x + (w - (bb[2] - bb[0])) / 2, cy), line, font=font, fill=fill)
            cy += lh + line_gap

    def rounded(x, y, w, h, fill, outline, label, icon=None, sub=None, radius=24, width=4):
        d.rounded_rectangle((x, y, x + w, y + h), radius=radius, fill=fill, outline=outline, width=width)
        if icon:
            d.text((x + 18, y + 18), icon, font=icon_font, fill=outline)
            tx = x + 64
            tw = w - 82
        else:
            tx = x + 16
            tw = w - 32
        text_center((tx, y + 11, tw, 28), label, node_font, ink)
        if sub:
            text_center((x + 18, y + 48, w - 36, h - 54), sub, small_font, ink)

    def cylinder(x, y, w, h, fill, outline, label, sub=None):
        d.rectangle((x, y + 28, x + w, y + h - 28), fill=fill, outline=outline, width=4)
        d.ellipse((x, y, x + w, y + 56), fill=fill, outline=outline, width=4)
        d.arc((x, y + h - 56, x + w, y + h), 0, 180, fill=outline, width=4)
        d.line((x, y + 28, x, y + h - 28), fill=outline, width=4)
        d.line((x + w, y + 28, x + w, y + h - 28), fill=outline, width=4)
        text_center((x + 18, y + 45, w - 36, 32), label, node_font, ink)
        if sub:
            text_center((x + 22, y + 86, w - 44, h - 105), sub, small_font, ink)

    def arrow(x1, y1, x2, y2, color=slate, width=5, label=None, curve=False):
        if curve:
            midx = (x1 + x2) / 2
            d.line((x1, y1, midx, y1, midx, y2, x2, y2), fill=color, width=width, joint="curve")
        else:
            d.line((x1, y1, x2, y2), fill=color, width=width)
        import math
        ang = math.atan2(y2 - y1, x2 - x1)
        r = 20
        p1 = (x2 - r * math.cos(ang - 0.45), y2 - r * math.sin(ang - 0.45))
        p2 = (x2 - r * math.cos(ang + 0.45), y2 - r * math.sin(ang + 0.45))
        d.polygon([(x2, y2), p1, p2], fill=color)
        if label:
            lx, ly = (x1 + x2) / 2, (y1 + y2) / 2 - 30
            bb = d.textbbox((0, 0), label, font=tiny_font)
            d.rounded_rectangle((lx - (bb[2] - bb[0]) / 2 - 10, ly - 6, lx + (bb[2] - bb[0]) / 2 + 10, ly + 24), radius=10, fill=(246, 250, 252), outline=grid, width=2)
            d.text((lx - (bb[2] - bb[0]) / 2, ly - 2), label, font=tiny_font, fill=ink)

    def lane(x, y, w, h, title, color):
        d.rounded_rectangle((x, y, x + w, y + h), radius=34, fill=(255, 255, 255), outline=color, width=4)
        d.rectangle((x, y, x + w, y + 56), fill=color)
        d.rounded_rectangle((x, y, x + w, y + h), radius=34, outline=color, width=4)
        d.text((x + 24, y + 12), title, font=h_font, fill=(255, 255, 255))

    d.text((70, 46), "AURA - AI Unified Research and Academic Tracker", font=title_font, fill=navy)
    d.text((72, 120), "End-to-end architecture for faculty academic-data collection, AI query understanding, SQL retrieval and automated accreditation reporting", font=small_font, fill=slate)
    d.line((70, 168, 2330, 168), fill=teal, width=7)

    # lanes
    lane(70, 230, 420, 930, "Users + Inputs", slate)
    lane(535, 230, 405, 930, "Presentation Layer", blue)
    lane(985, 230, 550, 930, "Application Services", teal)
    lane(1585, 230, 430, 930, "AI / NLP Layer", rose)
    lane(2060, 230, 270, 930, "Data + Output", gold)

    # user/input side
    rounded(105, 320, 345, 108, (241, 244, 246), slate, "Faculty / HOD / Admin", "role-based users", "●")
    rounded(105, 500, 345, 116, (255, 246, 230), gold, "Structured Sources", "Excel, department records,\nScholar, IEEE, Scopus", "▣")
    rounded(105, 700, 345, 110, (252, 236, 243), rose, "Document Sources", "certificates, PDFs,\naccreditation files", "◫")
    rounded(105, 905, 345, 102, (229, 246, 240), teal, "Institutional Reports", "NAAC, NBA, NIRF,\nfaculty evaluation", "◎")

    # presentation
    rounded(580, 330, 315, 120, (231, 241, 252), blue, "React Dashboard", "login, upload,\nquery, download", "▤")
    rounded(580, 560, 315, 110, (229, 246, 240), teal, "Workflow Views", "faculty summary,\nresults, status", "◧")
    rounded(580, 805, 315, 105, (241, 244, 246), slate, "Client Checks", "empty query +\nfile validation", "✓")

    # services
    rounded(1030, 320, 220, 105, (229, 246, 240), teal, "Auth + RBAC", "JWT/session\npermissions", "🔒")
    rounded(1280, 320, 220, 105, (229, 246, 240), teal, "REST API", "FastAPI endpoints\nrequest routing", "◇")
    rounded(1030, 505, 220, 130, (255, 246, 230), gold, "Ingestion", "Excel import\nmetadata extraction", "⇣")
    rounded(1280, 505, 220, 130, (255, 246, 230), gold, "Preprocess", "clean, normalize,\nmap schema", "⚙")
    rounded(1030, 730, 220, 130, (231, 241, 252), blue, "Query Engine", "intent + filters\nSQL execution", "⌕")
    rounded(1280, 730, 220, 130, (231, 241, 252), blue, "Report Service", "Excel/PDF\nreport templates", "▦")
    rounded(1155, 955, 245, 105, (252, 236, 243), rose, "Validation Layer", "SQL safety\noutput checks", "!")

    # AI
    rounded(1630, 330, 340, 110, (252, 236, 243), rose, "LLM Orchestration", "OpenAI + LangChain\nprompt pipeline", "◆")
    rounded(1630, 525, 340, 110, (252, 236, 243), rose, "Schema-Aware NL2SQL", "tables + columns\nas query context", "≋")
    rounded(1630, 720, 340, 110, (252, 236, 243), rose, "Query Refinement", "clarify intent\nreduce SQL errors", "↺")
    rounded(1630, 915, 340, 105, (252, 236, 243), rose, "Future RAG + OCR", "PDF extraction\nexternal APIs", "◇")

    # data/output
    cylinder(2090, 325, 210, 205, (231, 241, 252), blue, "Academic DB", "faculty, papers,\nFDPs, patents,\ncertifications")
    cylinder(2090, 620, 210, 155, (241, 244, 246), slate, "Audit Logs", "access trace\nquery status")
    rounded(2090, 880, 210, 140, (229, 246, 240), teal, "Outputs", "dashboard\nExcel / PDF\nanalytics", "⬇")

    # main flow arrows
    arrow(450, 374, 580, 390, slate, label="access")
    arrow(450, 558, 580, 390, gold, label="upload")
    arrow(895, 390, 1030, 372, blue, label="API")
    arrow(895, 858, 1030, 795, slate, label="checks")
    arrow(1250, 570, 1280, 570, gold)
    arrow(1500, 570, 1630, 580, rose, label="schema")
    arrow(1500, 795, 1630, 775, rose, label="NL query")
    arrow(1970, 580, 2090, 428, rose, label="SQL")
    arrow(1500, 795, 2090, 428, teal, label="execute")
    arrow(2195, 530, 2195, 620, slate, label="log")
    arrow(2195, 775, 2195, 880, teal, label="report")
    arrow(2090, 950, 1500, 795, teal, label="results", curve=True)
    arrow(1280, 795, 895, 615, blue, label="response", curve=True)

    # security strip
    d.rounded_rectangle((590, 1195, 2285, 1265), radius=24, fill=(232, 242, 239), outline=teal, width=3)
    d.text((625, 1214), "Cross-cutting controls: authentication | role-based access | SQL validation | data normalization | audit logging | report-template consistency | secure storage", font=small_font, fill=navy)

    img.save(path)


def create_diagram_pngs():
    colors = {
        "blue": ((231, 241, 252), (24, 95, 165)),
        "green": ((229, 246, 240), (15, 110, 86)),
        "gold": ((255, 246, 230), (186, 117, 23)),
        "rose": ((252, 236, 243), (153, 53, 86)),
        "gray": ((241, 244, 246), (92, 104, 120)),
    }
    arch_nodes = {
        "sources": (70, 180, 290, 120, "Academic Data\nExcel / Scholar / IEEE", *colors["gold"]),
        "ui": (435, 180, 260, 120, "React Web\nDashboard", *colors["blue"]),
        "api": (770, 180, 270, 120, "FastAPI\nBackend", *colors["green"]),
        "ai": (1115, 160, 320, 150, "AI/NLP Layer\nOpenAI + LangChain", *colors["rose"]),
        "db": (770, 430, 270, 120, "PostgreSQL /\nMySQL Database", *colors["blue"]),
        "report": (1115, 430, 320, 120, "Report Engine\nExcel / PDF", *colors["green"]),
        "users": (435, 430, 260, 120, "Faculty / HOD\nAdmin Users", *colors["gray"]),
    }
    draw_png_diagram(DESKTOP / "AURA_system_architecture.png", "AURA System Architecture", arch_nodes,
                     [("sources", "ui"), ("ui", "api"), ("api", "ai"), ("ai", "db"), ("api", "db"), ("db", "report"), ("report", "users"), ("users", "ui")])
    draw_architecture_png(DESKTOP / "AURA_system_architecture_detailed.png")
    dfd_nodes = {
        "user": (80, 220, 250, 100, "User Query /\nData Upload", *colors["gray"]),
        "front": (420, 220, 250, 100, "Frontend\nValidation", *colors["blue"]),
        "prep": (760, 220, 260, 100, "Preprocess +\nSchema Mapping", *colors["gold"]),
        "llm": (1100, 220, 300, 100, "LLM SQL\nGeneration", *colors["rose"]),
        "db": (760, 470, 260, 100, "Faculty Academic\nDatabase", *colors["green"]),
        "out": (1100, 470, 300, 100, "Results +\nReport Download", *colors["blue"]),
    }
    draw_png_diagram(DESKTOP / "AURA_dataflow_diagram.png", "AURA Data Flow Diagram", dfd_nodes,
                     [("user", "front"), ("front", "prep"), ("prep", "llm"), ("llm", "db"), ("db", "out"), ("out", "user")])
    use_nodes = {
        "fac": (80, 220, 210, 95, "Faculty /\nHOD", *colors["gray"]),
        "admin": (80, 500, 210, 95, "Admin", *colors["gray"]),
        "login": (430, 160, 250, 80, "Register / Login", *colors["blue"]),
        "upload": (430, 280, 250, 80, "Upload Dataset", *colors["gold"]),
        "query": (430, 400, 250, 80, "Ask NL Query", *colors["rose"]),
        "report": (790, 280, 280, 80, "Generate Report", *colors["green"]),
        "manage": (790, 400, 280, 80, "Manage Users /\nValidate Data", *colors["blue"]),
        "download": (1120, 340, 280, 80, "Download Excel /\nPDF Output", *colors["green"]),
    }
    draw_png_diagram(DESKTOP / "AURA_usecase_diagram.png", "AURA Use Case Diagram", use_nodes,
                     [("fac", "login"), ("fac", "upload"), ("fac", "query"), ("query", "report"), ("report", "download"), ("admin", "manage"), ("admin", "report")])
    draw_sequence_png(DESKTOP / "AURA_sequence_diagram.png")


def add_architecture(slide):
    add_header(slide, "SYSTEM ARCHITECTURE", 10)
    img = DESKTOP / "AURA_system_architecture_detailed.png"
    slide.shapes.add_picture(str(img), Inches(0.45), Inches(0.92), width=Inches(12.45), height=Inches(5.9))


def add_dataflow(slide):
    add_header(slide, "DATAFLOW DIAGRAM", 12)
    labels = [
        ("User query /\ndata upload", 0.75, 1.7, "gray"),
        ("Frontend\nvalidation", 3.0, 1.7, "blue"),
        ("Preprocessing +\nschema mapping", 5.35, 1.7, "gold"),
        ("LLM converts\nNL to SQL", 7.95, 1.7, "rose"),
        ("Faculty academic\ndatabase", 5.35, 4.25, "green"),
        ("Result view +\nExcel/PDF report", 8.05, 4.25, "blue"),
    ]
    color_map = {
        "gray": (RGBColor(241, 244, 246), SLATE),
        "blue": (RGBColor(231, 241, 252), RGBColor(24, 95, 165)),
        "gold": (RGBColor(255, 246, 230), GOLD),
        "rose": (RGBColor(252, 236, 243), ROSE),
        "green": (RGBColor(229, 246, 240), TEAL),
    }
    boxes = []
    for text, x, y, c in labels:
        boxes.append(add_box(slide, Inches(x), Inches(y), Inches(2.0), Inches(0.82), text, *color_map[c], size=12))
    coords = [(1, 2), (2, 3), (3, 4), (4, 5), (5, 6), (6, 1)]
    for i, j in coords:
        a, b = boxes[i-1], boxes[j-1]
        add_arrow(slide, a.left + a.width, a.top + a.height/2, b.left, b.top + b.height/2)
    add_textbox(slide, Inches(0.8), Inches(5.85), Inches(11.5), Inches(0.55),
                "Flow: collect clean data, interpret intent, generate safe SQL, retrieve records, validate output, and export a structured institutional report.",
                16, False, INK)


def add_usecase(slide):
    add_header(slide, "USE CASE DIAGRAM", 13)
    boundary = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(3.0), Inches(1.25), Inches(7.1), Inches(5.25))
    boundary.fill.solid(); boundary.fill.fore_color.rgb = RGBColor(250, 252, 253)
    boundary.line.color.rgb = SLATE; boundary.line.width = Pt(1.2)
    add_textbox(slide, Inches(3.25), Inches(1.38), Inches(2.2), Inches(0.25), "AURA System", 14, True, NAVY)
    add_box(slide, Inches(0.75), Inches(2.1), Inches(1.6), Inches(0.6), "Faculty / HOD", RGBColor(241,244,246), SLATE, 12)
    add_box(slide, Inches(0.75), Inches(4.9), Inches(1.6), Inches(0.6), "Admin", RGBColor(241,244,246), SLATE, 12)
    usecases = [
        ("Register / Login", 3.65, 1.95),
        ("Upload academic data", 3.65, 3.05),
        ("Ask natural-language query", 3.65, 4.15),
        ("Generate Excel / PDF report", 6.65, 3.05),
        ("Manage users + validate data", 6.65, 4.15),
        ("View dashboard analytics", 6.65, 5.25),
    ]
    for text, x, y in usecases:
        shp = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL, Inches(x), Inches(y), Inches(2.45), Inches(0.58))
        shp.fill.solid(); shp.fill.fore_color.rgb = RGBColor(229, 246, 240)
        shp.line.color.rgb = TEAL
        shp.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
        set_text(shp, text, 11, True, INK, PP_ALIGN.CENTER)
    for y in [2.25, 3.35, 4.45]:
        add_arrow(slide, Inches(2.35), Inches(2.4), Inches(3.65), Inches(y), SLATE, 1.3)
    for y in [4.45, 5.55]:
        add_arrow(slide, Inches(2.35), Inches(5.2), Inches(6.65), Inches(y), SLATE, 1.3)


def add_sequence(slide):
    add_header(slide, "SEQUENCE DIAGRAM", 14)
    names = [("User", 0.75), ("React UI", 3.0), ("FastAPI", 5.15), ("AI/LLM Layer", 7.4), ("DB + Reports", 9.85)]
    tops = {}
    for name, x in names:
        add_box(slide, Inches(x), Inches(1.2), Inches(1.7), Inches(0.45), name, RGBColor(231,241,252), RGBColor(24,95,165), 10)
        line = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x+0.85), Inches(1.7), Inches(x+0.85), Inches(6.35))
        line.line.color.rgb = SLATE; line.line.width = Pt(1)
        tops[name] = Inches(x+0.85)
    steps = [
        ("Submit query / upload data", "User", "React UI", 2.0),
        ("POST request", "React UI", "FastAPI", 2.55),
        ("preprocess + prompt", "FastAPI", "AI/LLM Layer", 3.1),
        ("schema-aware SQL", "AI/LLM Layer", "FastAPI", 3.65),
        ("execute validated SQL", "FastAPI", "DB + Reports", 4.2),
        ("records + report file", "DB + Reports", "FastAPI", 4.75),
        ("response + download link", "FastAPI", "React UI", 5.3),
        ("display result", "React UI", "User", 5.85),
    ]
    for label, src, dst, y in steps:
        add_arrow(slide, tops[src], Inches(y), tops[dst], Inches(y), TEAL if src != "AI/LLM Layer" else ROSE, 1.5)
        add_textbox(slide, min(tops[src], tops[dst]) + Inches(0.05), Inches(y-0.2), abs(tops[dst]-tops[src]) - Inches(0.1), Inches(0.18), label, 8.8, False, INK, PP_ALIGN.CENTER)


def build_deck():
    create_diagram_pngs()
    prs = Presentation(str(SRC))
    while len(prs.slides) < 18:
        prs.slides.add_slide(prs.slide_layouts[6])

    slide_data = {
        1: ("AURA: AI Unified Research and Academic Tracker for Faculty", [
            "Project Phase-II Report [22IS81]",
            "Presented by: Santrupthi R B, Shreya Hunnur, Siddu H G, Chethan G",
            "Guide: Mr. Yogesh B S, Assistant Professor, Dept. of ISE, DSCE",
        ]),
        2: ("CONTENTS", ["Abstract", "Introduction", "Literature Survey", "Research Gap", "Problem Statement", "Hardware & Software Requirements", "System Design & Methodology", "System Architecture", "Data Flow Diagram", "Use Case Diagram", "Sequence Diagram", "Results", "References", "Project Outcomes"]),
        3: ("ABSTRACT", ["AURA is an AI-based academic management platform designed to track faculty activities such as publications, FDPs, conferences, patents, certifications, research work and accreditation records.", "The system replaces manual searching through spreadsheets and files with a natural-language interface where users can ask queries like “show IEEE publications for 2024” or “generate FDP report”.", "Natural Language Processing and Large Language Models understand user intent, map it with the database schema and generate validated SQL queries automatically.", "Cleaned faculty data from Excel sheets, institutional records and academic platforms is stored in a centralized relational database for consistent retrieval.", "AURA generates structured Excel/PDF reports through a user-friendly dashboard, reducing manual effort, missing records, duplicate entries and reporting errors."]),
        4: ("INTRODUCTION", ["Educational institutions maintain large volumes of faculty academic data for publications, journals, conferences, FDPs, patents, certifications, projects and professional activities.", "In many departments this information is scattered across spreadsheets, Google Scholar profiles, ResearchGate, IEEE Xplore, institutional portals and physical documents.", "Accreditation activities such as NAAC, NBA and NIRF require accurate data in a fixed format within a limited time, making manual preparation stressful and repetitive.", "AURA introduces an AI-driven workflow where faculty or administrators can upload records, ask natural-language queries, retrieve filtered data and generate reports instantly.", "The project combines centralized data management, schema-aware SQL generation, validation, dashboard visualization and automated report generation into one practical system."]),
        5: ("LITERATURE SURVEY", ["Generative AI in academia shows that LLMs can automate repetitive research and administrative workflows while improving productivity in educational institutions.", "Text-to-SQL research demonstrates that modern LLMs can translate natural language into database queries, allowing non-technical users to access structured records.", "Faculty publication aggregation studies highlight the need for automated collection and metadata extraction from academic web pages and institutional sources.", "Retrieval-based SQL frameworks such as RB-SQL show that database-schema context improves SQL accuracy and reduces wrong or hallucinated query generation.", "Survey papers on LLM database interfaces support prompt engineering, schema linking, query validation and natural-language interfaces as key parts of reliable AI systems.", "These studies form the base for AURA’s design: AI query understanding, schema-aware SQL, academic data retrieval and automated report generation."]),
        6: ("RESEARCH GAP", ["Most existing academic systems focus on either data storage or report preparation, but they do not provide an integrated end-to-end AI workflow.", "Faculty records are still commonly maintained manually in disconnected spreadsheets, leading to repeated data entry and inconsistent formatting.", "Many tools require SQL knowledge or technical support, which limits direct use by faculty members, HODs and accreditation coordinators.", "Existing systems do not handle natural-language academic queries such as department-wise publication filters, year-wise FDP reports or conference summaries effectively.", "Automated generation of institutional report formats for NAAC, NBA, NIRF and internal reviews is still limited.", "There is a clear need for a centralized, AI-enabled, secure and scalable platform that can collect, validate, query and report faculty academic data."]),
        7: ("PROBLEM STATEMENT", ["Faculty academic records are distributed across multiple files, online platforms and department-level documents, making retrieval slow and unreliable.", "Manual record maintenance causes duplicate entries, missing publication details, inconsistent author names, incorrect years and format mismatches.", "During accreditation or faculty evaluation, administrators must manually filter records, verify data and prepare reports under strict deadlines.", "Non-technical users cannot easily query databases because they need SQL knowledge or backend support.", "The absence of automated validation and report generation increases human error and reduces confidence in institutional data.", "AURA solves this by using natural-language query processing, schema-aware SQL generation, centralized storage, role-based access and automated Excel/PDF reporting."]),
        8: ("HARDWARE & SOFTWARE REQUIREMENTS", ["Hardware: Intel Core i5 / AMD Ryzen 5 or higher processor, minimum 8 GB RAM, 16 GB recommended for smoother AI and backend execution.", "Storage: Minimum 256 GB SSD preferred for project files, datasets, generated reports and database storage.", "Network: Stable internet connection for API communication, model access, GitHub collaboration and future external academic-platform integrations.", "Frontend: React.js, HTML5, CSS3 and JavaScript for the interactive web dashboard, query form, upload workflow and report download interface.", "Backend: Python with FastAPI for REST APIs, authentication flow, query orchestration, validation and report-generation endpoints.", "AI/NLP: OpenAI SDK, LangChain, LangGraph, Transformers and prompt engineering for natural-language understanding and SQL generation.", "Database & tools: PostgreSQL/MySQL, SQLAlchemy, OpenPyXL, Tableau, Postman, VS Code, Git and GitHub."]),
        9: ("SYSTEM DESIGN & METHODOLOGY", ["Data Collection: Faculty records are collected from Excel sheets, department files, institutional databases and academic platforms such as Google Scholar, IEEE Xplore and ResearchGate.", "Preprocessing: Raw records are cleaned, normalized, de-duplicated and validated to handle missing fields, spelling variations and inconsistent formats.", "Schema Mapping: Cleaned records are mapped into relational tables such as faculty, publications, FDPs, conferences, patents, certifications and reports.", "AI Query Processing: User queries are processed using NLP and LLM workflows to identify intent, entities, filters, years, departments and report type.", "SQL Generation: The system generates schema-aware SQL and validates table names, columns, conditions and query safety before execution.", "Report Generation: Retrieved results are converted into structured Excel/PDF outputs and displayed through the dashboard for download and review."]),
        11: ("PROPOSED ALGORITHM AND TECHNIQUES", ["Step 1: Authenticate the user and identify the role as Faculty, HOD or Admin.", "Step 2: Accept either dataset upload or natural-language query from the web dashboard.", "Step 3: Clean uploaded data by removing duplicates, standardizing fields and validating mandatory values.", "Step 4: Extract query intent, entities and filters such as faculty name, department, year, publication type or accreditation category.", "Step 5: Provide schema context to the LLM and generate an SQL query mapped to the correct database tables.", "Step 6: Validate generated SQL for correctness, safety and schema compatibility before execution.", "Step 7: Execute the query, retrieve academic records and format the response for dashboard display.", "Step 8: Generate Excel/PDF report using predefined institutional templates and allow download from the interface."]),
        15: ("RESULTS", ["AURA successfully supports academic queries for publications, conferences, FDPs, patents, certifications, journals and accreditation-related records.", "The dashboard provides a single interface for dataset upload, natural-language query entry, workflow tracking, result visualization and report download.", "Schema-aware prompting improves the correctness of generated SQL by grounding queries in the actual database structure.", "Validation steps reduce unsafe queries, wrong column references, duplicate outputs and missing report fields.", "Automated Excel/PDF report generation reduces the time required for manual spreadsheet filtering and formatting.", "Centralized storage improves consistency across department records and makes faculty academic data easier to retrieve during institutional reviews.", "The system demonstrates practical use of AI, NLP and database automation for academic administration."]),
        16: ("REFERENCES", ["A. M. Hanfi, O. A. Al-Shariff and M. S. Ahmed, “Generative AI in Academia: A Comprehensive Review of Applications and Implications for Research,” International Journal of Engineering and Applied Sciences, 2025.", "A. Mohammad Jafari, A. S. Maida and R. Gottumukkala, “From Natural Language to SQL: Review of LLM-Based Text-to-SQL Systems,” arXiv preprint, 2024.", "N. Jahn, M. Lösch and W. Horstmann, “Automatic Aggregation of Faculty Publications from Personal Web Pages,” Code4Lib Journal, 2010.", "Z. Wu et al., “RB-SQL: A Retrieval-Based LLM Framework for Text-to-SQL,” arXiv preprint, 2024.", "Z. Hong et al., “Next-Generation Database Interfaces: A Survey of LLM-Based Text-to-SQL Applications,” IEEE Transactions on Knowledge and Data Engineering, 2025.", "A. B. Kanburoğlu and F. B. Tek, “Text-to-SQL: A Methodical Review of Challenges and Models,” Turkish Journal of Electrical Engineering and Computer Sciences, 2024."]),
        17: ("PROJECT OUTCOMES", ["A centralized academic tracker for faculty publications, FDPs, conferences, patents, certifications, journals and accreditation records.", "Natural-language access to faculty data, reducing the need for manual SQL writing or repeated spreadsheet searches.", "Automated generation of structured Excel/PDF reports suitable for department reviews, accreditation support and faculty evaluation.", "Improved data quality through preprocessing, duplicate removal, schema mapping and validation before storage and reporting.", "Reduced manual workload for faculty members, HODs and administrators during NAAC, NBA, NIRF and internal academic reviews.", "Secure and scalable architecture with scope for role-based access control, audit logs, OCR-based certificate extraction and external academic API integration.", "Future enhancement: RAG-based query refinement, Google Scholar/Scopus/ORCID integration, advanced analytics dashboard and stronger AI safety checks."]),
        18: ("THANK YOU", ["AURA: AI Unified Research and Academic Tracker for Faculty", "Department of Information Science & Engineering", "Dayananda Sagar College of Engineering"]),
    }

    for idx, slide in enumerate(prs.slides, 1):
        clear_slide(slide)
        bg = slide.background.fill
        bg.solid()
        bg.fore_color.rgb = WHITE
        if idx == 1:
            slide.background.fill.fore_color.rgb = RGBColor(247, 250, 251)
            add_textbox(slide, Inches(0.65), Inches(0.35), Inches(7.5), Inches(0.3), "DEPARTMENT OF INFORMATION SCIENCE & ENGINEERING", 13, True, TEAL)
            add_textbox(slide, Inches(0.65), Inches(1.35), Inches(11.0), Inches(1.4), slide_data[idx][0], 34, True, NAVY)
            add_textbox(slide, Inches(0.7), Inches(3.05), Inches(10.9), Inches(1.3), "\n".join(slide_data[idx][1]), 15, False, INK)
            add_box(slide, Inches(0.7), Inches(5.65), Inches(3.6), Inches(0.55), "Project Phase-II | 2025-26", RGBColor(229,246,240), TEAL, 14)
            add_textbox(slide, Inches(0.65), Inches(6.95), Inches(8), Inches(0.2), "Visvesvaraya Technological University", 9, False, SLATE)
            continue
        if idx in (10, 12, 13, 14):
            if idx == 10:
                add_architecture(slide)
            elif idx == 12:
                add_dataflow(slide)
            elif idx == 13:
                add_usecase(slide)
            elif idx == 14:
                add_sequence(slide)
            continue
        title, bullets = slide_data.get(idx, ("AURA", []))
        add_header(slide, title, idx)
        if idx == 2:
            cols = [bullets[:5], bullets[5:10], bullets[10:]]
            for c, items in enumerate(cols):
                add_bullets(slide, items, Inches(0.85 + c*4.1), Inches(1.35), Inches(3.55), Inches(5.4), 17)
        elif idx == 18:
            add_textbox(slide, Inches(1.1), Inches(2.0), Inches(11.2), Inches(1.1), "THANK YOU", 54, True, NAVY, PP_ALIGN.CENTER)
            add_textbox(slide, Inches(1.8), Inches(3.45), Inches(9.8), Inches(1.0), "\n".join(bullets), 20, False, TEAL, PP_ALIGN.CENTER)
        else:
            add_bullets(slide, bullets, Inches(0.95), Inches(1.35), Inches(11.2), Inches(5.4), 19 if idx not in (8,16) else 16)

    prs.save(str(OUT))


if __name__ == "__main__":
    build_deck()
    print(OUT)
    for name in [
        "AURA_system_architecture.png",
        "AURA_dataflow_diagram.png",
        "AURA_usecase_diagram.png",
        "AURA_sequence_diagram.png",
    ]:
        print(DESKTOP / name)
